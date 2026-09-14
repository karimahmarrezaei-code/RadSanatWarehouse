# -*- coding: utf-8 -*-
"""Report Export Service - نسخه نهایی با فاصله‌گذاری و واحد ریال"""

from pathlib import Path
from typing import Any, Dict, List

import arabic_reshaper
import jdatetime
from bidi.algorithm import get_display

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OPERATION_LABELS = {
    'INBOUND_RECEIPT': 'خرید / رسید ورودی',
    'INBOUND_FREIGHT': 'کرایه حمل ورودی',
    'OUTBOUND_ISSUE': 'فروش / حواله خروج',
    'OUTBOUND_FREIGHT': 'کرایه حمل خروجی',
}
ROW_TYPE_LABELS = {'FINANCE_DOC': 'سند مالی', 'PAYMENT': 'تسویه'}


# ================================================================
# 🔤 ثبت فونت فارسی
# ================================================================
_FONT_CANDIDATES = [
    Path('fonts/Vazirmatn-Regular.ttf'),
    Path('fonts/Vazir.ttf'),
    Path('fonts/IRANSans.ttf'),
    Path('C:/Windows/Fonts/tahoma.ttf'),
    Path('/usr/share/fonts/truetype/vazirmatn/Vazirmatn-Regular.ttf'),
]

FA_FONT = 'Helvetica'
for _p in _FONT_CANDIDATES:
    if _p.exists():
        try:
            pdfmetrics.registerFont(TTFont('FaFont', str(_p)))
            FA_FONT = 'FaFont'
            break
        except Exception:
            continue


# ================================================================
# 🔧 توابع کمکی
# ================================================================
def fa(text: Any) -> str:
    """شکل‌دهی متن فارسی برای PDF"""
    if text is None:
        return ''
    text = str(text).strip()
    if not text:
        return ''
    try:
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def _fmt(value: Any) -> str:
    try:
        return f'{int(value):,}'
    except Exception:
        return str(value) if value is not None else '-'


def _money(value: Any) -> str:
    """مبلغ با جداکننده + واحد ریال"""
    try:
        return f'{int(value):,} ریال'
    except Exception:
        return '-'


def _int(value: Any) -> int:
    try:
        return int(value or 0)
    except Exception:
        return 0


def _fa_date(value: Any) -> str:
    """تبدیل تاریخ میلادی (yyyy-mm-dd) به شمسی"""
    if not value:
        return '-'
    text = str(value).strip()
    try:
        parts = text.split(' ')[0].split('-')
        if len(parts) == 3:
            g = jdatetime.date.fromgregorian(
                year=int(parts[0]), month=int(parts[1]), day=int(parts[2]))
            return g.strftime('%Y/%m/%d')
    except Exception:
        pass
    return text


def _today_fa() -> str:
    return jdatetime.date.today().strftime('%Y/%m/%d')


# ================================================================
# بلوک‌های PDF
# ================================================================
def _text_block(text, size=10, color='#0f172a', align='CENTER') -> Table:
    t = Table([[fa(text)]])
    t.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), FA_FONT),
        ('FONTSIZE', (0, 0), (-1, -1), size),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor(color)),
        ('ALIGN', (0, 0), (-1, -1), align),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    return t


def _summary_table(pairs: List, usable_width) -> Table:
    rows = []
    for i in range(0, len(pairs), 2):
        row = []
        for j in range(2):
            if i + j < len(pairs):
                label, value = pairs[i + j]
                row += [label, value]
            else:
                row += ['', '']
        rows.append(row)

    data = [[fa(c) for c in r] for r in rows]
    w = usable_width
    t = Table(data, colWidths=[w * 0.15, w * 0.35, w * 0.15, w * 0.35])
    t.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), FA_FONT),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#64748b')),
        ('TEXTCOLOR', (2, 0), (2, -1), colors.HexColor('#64748b')),
        ('TEXTCOLOR', (1, 0), (1, -1), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (3, 0), (3, -1), colors.HexColor('#0f172a')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f1f5f9')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#f1f5f9')),
        ('BOX', (0, 0), (-1, -1), 0.6, colors.HexColor('#94a3b8')),
        ('INNERGRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
    ]))
    return t


def _data_table(headers, rows, widths, font_size=8.5) -> Table:
    data = [[fa(h) for h in headers]]
    for r in rows:
        data.append([fa(c) for c in r])
    t = Table(data, colWidths=widths, repeatRows=1, hAlign='CENTER')
    t.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), FA_FONT),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), font_size),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#94a3b8')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#eef2f7')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
    ]))
    return t


def _footer_factory(app_name: str):
    def _draw(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#cbd5e1'))
        canvas.setLineWidth(0.6)
        canvas.line(1.2 * cm, 1.1 * cm, doc.pagesize[0] - 1.2 * cm, 1.1 * cm)
        canvas.setFont(FA_FONT, 8)
        canvas.setFillColor(colors.HexColor('#64748b'))
        canvas.drawRightString(doc.pagesize[0] - 1.2 * cm, 0.65 * cm, fa(f'صفحه {doc.page}'))
        canvas.drawCentredString(doc.pagesize[0] / 2, 0.65 * cm, fa(f'تاریخ گزارش:   {_today_fa()}'))
        canvas.drawString(1.2 * cm, 0.65 * cm, fa(app_name))
        canvas.restoreState()
    return _draw


def _usable_width():
    return landscape(A4)[0] - 2.4 * cm


def _build_pdf(path, story, app_name='سیستم مدیریت پالت انبار') -> str:
    doc = SimpleDocTemplate(
        str(path), pagesize=landscape(A4),
        leftMargin=1.2 * cm, rightMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.4 * cm,
    )
    footer = _footer_factory(app_name)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return str(path)


# ================================================================
# 📗 توابع Excel
# ================================================================
def _write_excel(path, sheets: List[Dict]) -> str:
    wb = Workbook()
    wb.remove(wb.active)

    hfill = PatternFill('solid', fgColor='1E40AF')
    hfont = Font(bold=True, color='FFFFFF', size=10, name='Tahoma')
    bfont = Font(bold=True, size=10, name='Tahoma')
    tfont = Font(bold=True, size=14, color='0F172A', name='Tahoma')
    sfont = Font(size=9, color='64748B', name='Tahoma')
    nfont = Font(size=10, name='Tahoma')
    thin = Side(style='thin', color='CBD5E1')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt_fill = PatternFill('solid', fgColor='EEF2F7')

    for spec in sheets:
        ws = wb.create_sheet(str(spec['name'])[:31])
        ws.sheet_view.rightToLeft = True
        row_idx = 1

        if spec.get('title'):
            ws.cell(row=row_idx, column=1, value=spec['title']).font = tfont
            row_idx += 1
        if spec.get('subtitle'):
            ws.cell(row=row_idx, column=1, value=spec['subtitle']).font = sfont
            row_idx += 1

        for label, value in spec.get('summary', []):
            ws.cell(row=row_idx, column=1, value=label).font = bfont
            c2 = ws.cell(row=row_idx, column=2, value=value)
            c2.font = nfont
            if isinstance(value, (int, float)):
                c2.number_format = '#,##0'
            row_idx += 1
        if spec.get('summary'):
            row_idx += 1

        header_row = row_idx
        for ci, h in enumerate(spec['headers'], start=1):
            c = ws.cell(row=header_row, column=ci, value=h)
            c.fill = hfill
            c.font = hfont
            c.border = border
            c.alignment = Alignment(horizontal='center', vertical='center')

        for ri, r in enumerate(spec['rows']):
            for ci, val in enumerate(r, start=1):
                c = ws.cell(row=header_row + 1 + ri, column=ci, value=val)
                c.border = border
                c.font = nfont
                c.alignment = Alignment(horizontal='right', vertical='center')
                if isinstance(val, (int, float)):
                    c.number_format = '#,##0'
                if ri % 2 == 1:
                    c.fill = alt_fill

        for ci, w in enumerate(spec.get('widths', []), start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w

        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    wb.save(str(path))
    return str(path)


# ================================================================
# 📑 اسناد مالی
# ================================================================
def export_financial_summary_to_excel(report, file_path) -> str:
    s = report['summary']
    rows = []
    for it in report['items']:
        date_label = it.get('finance_date_label') or _fa_date(it.get('finance_date')) or '-'
        rows.append([
            it['finance_no'], date_label, it['operation_type_label'], it['direction_label'],
            it.get('counterparty_name') or '-', it['reference_label'],
            _int(it.get('total_amount')), _int(it.get('settled_amount')),
            _int(it.get('remaining_amount')), it['status_label'],
        ])
    return _write_excel(file_path, [{
        'name': 'اسناد مالی',
        'title': 'گزارش اسناد مالی',
        'subtitle': f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال',
        'summary': [
            ('تعداد کل اسناد', _int(s['total_docs'])),
            ('اسناد باز', _int(s['open_docs'])),
            ('اسناد تسویه', _int(s['settled_docs'])),
            ('کل دریافتنی', _int(s['receivable_total'])),
            ('مانده دریافتنی', _int(s['receivable_balance'])),
            ('کل پرداختنی', _int(s['payable_total'])),
            ('مانده پرداختنی', _int(s['payable_balance'])),
        ],
        'headers': ['شماره سند', 'تاریخ', 'نوع عملیات', 'جهت', 'طرف حساب', 'مرجع عملیات', 'مبلغ کل (ریال)', 'تسویه شده (ریال)', 'مانده (ریال)', 'وضعیت'],
        'rows': rows,
        'widths': [14, 14, 22, 12, 24, 16, 18, 18, 18, 10],
    }])


def export_financial_summary_to_pdf(report, file_path) -> str:
    s = report['summary']
    w = _usable_width()
    story = [
        _text_block('گزارش اسناد مالی', 16),
        _text_block(f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال', 9, '#64748b'),
        Spacer(1, 4),
        _summary_table([
            ('تعداد کل اسناد', _fmt(s['total_docs'])),
            ('اسناد باز', _fmt(s['open_docs'])),
            ('اسناد تسویه', _fmt(s['settled_docs'])),
            ('کل دریافتنی', _money(s['receivable_total'])),
            ('مانده دریافتنی', _money(s['receivable_balance'])),
            ('کل پرداختنی', _money(s['payable_total'])),
            ('مانده پرداختنی', _money(s['payable_balance'])),
        ], w),
        Spacer(1, 8),
    ]
    headers = ['شماره سند', 'تاریخ', 'نوع عملیات', 'جهت', 'طرف حساب', 'مرجع', 'مبلغ کل', 'تسویه شده', 'مانده', 'وضعیت']
    rows = []
    for it in report['items']:
        date_label = it.get('finance_date_label') or _fa_date(it.get('finance_date')) or '-'
        rows.append([
            it['finance_no'], date_label, it['operation_type_label'], it['direction_label'],
            it.get('counterparty_name') or '-', it['reference_label'],
            _money(it.get('total_amount')), _money(it.get('settled_amount')),
            _money(it.get('remaining_amount')), it['status_label'],
        ])
    widths = [2.3, 2.3, 2.9, 2.0, 3.4, 2.2, 2.9, 2.9, 2.7, 1.5]
    story.append(_data_table(headers, rows, [x * cm for x in widths], font_size=8))
    return _build_pdf(file_path, story)


# ================================================================
# 👥 گردش حساب شخص
# ================================================================
def export_person_statement_to_excel(statement, file_path) -> str:
    person = statement['person']
    summary = statement['summary']
    full_name = f"{person['first_name']} {person['last_name']}".strip()
    rows = []
    for d in statement['rows']:
        rows.append([
            _fa_date(d['event_date']), ROW_TYPE_LABELS.get(d['row_type'], d['row_type']),
            d['reference_no'], OPERATION_LABELS.get(d['operation_type'], d['operation_type']),
            d.get('description') or '-',
            _int(d['debit_amount']), _int(d['credit_amount']), _int(d['running_balance']),
        ])
    return _write_excel(file_path, [{
        'name': 'گردش حساب',
        'title': f'گردش حساب:   {full_name}',
        'subtitle': f'موبایل:   {person.get("mobile") or "-"}   |   تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال',
        'summary': [
            ('کل دریافتنی', _int(summary['receivable_total'])),
            ('کل پرداختنی', _int(summary['payable_total'])),
            ('خالص مانده', _int(summary['net_balance'])),
        ],
        'headers': ['تاریخ', 'نوع ردیف', 'شماره سند', 'نوع عملیات', 'شرح', 'بدهکار (ریال)', 'بستانکار (ریال)', 'مانده (ریال)'],
        'rows': rows,
        'widths': [12, 12, 14, 20, 30, 18, 18, 18],
    }])


def export_person_statement_to_pdf(statement, file_path) -> str:
    person = statement['person']
    summary = statement['summary']
    full_name = f"{person['first_name']} {person['last_name']}".strip()
    w = _usable_width()
    story = [
        _text_block(f'گردش حساب:   {full_name}', 16),
        _text_block(f'موبایل:   {person.get("mobile") or "-"}   |   تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال', 9, '#64748b'),
        Spacer(1, 4),
        _summary_table([
            ('کل دریافتنی', _money(summary['receivable_total'])),
            ('کل پرداختنی', _money(summary['payable_total'])),
            ('خالص مانده', _money(summary['net_balance'])),
        ], w),
        Spacer(1, 8),
    ]
    headers = ['تاریخ', 'نوع ردیف', 'شماره سند', 'نوع عملیات', 'شرح', 'بدهکار', 'بستانکار', 'مانده']
    rows = []
    for d in statement['rows']:
        running = _int(d['running_balance'])
        running_text = f'{_money(abs(running))}  {"بدهکار" if running >= 0 else "بستانکار"}'
        rows.append([
            _fa_date(d['event_date']), ROW_TYPE_LABELS.get(d['row_type'], d['row_type']),
            d['reference_no'], OPERATION_LABELS.get(d['operation_type'], d['operation_type']),
            d.get('description') or '-',
            _money(d['debit_amount']) if _int(d['debit_amount']) else '-',
            _money(d['credit_amount']) if _int(d['credit_amount']) else '-',
            running_text,
        ])
    widths = [2.4, 2.2, 2.8, 3.3, 5.5, 3.0, 3.0, 4.9]
    story.append(_data_table(headers, rows, [x * cm for x in widths], font_size=8))
    return _build_pdf(file_path, story)


# ================================================================
# 🏭 ارزش ریالی انبار
# ================================================================
def export_stock_value_to_excel(report, file_path) -> str:
    s = report['summary']
    rows = []
    for i, it in enumerate(report['items'], 1):
        rows.append([
            i, it['pallet_code'], it['pallet_name'], it['material_type'], it['dimensions'],
            _int(it['total_in_qty']), _int(it['total_out_qty']), _int(it['current_qty']),
            _int(it['total_in_value']), _int(it['total_out_value']), _int(it['current_value']),
            it.get('last_transaction_date') or '-',
        ])
    return _write_excel(file_path, [{
        'name': 'ارزش ریالی',
        'title': f"گزارش موجودی و ارزش ریالی انبار:   {report['warehouse']['name']}",
        'subtitle': f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال',
        'summary': [
            ('کد انبار', report['warehouse']['code']),
            ('تعداد کل موجودی', _int(s['total_qty'])),
            ('ارزش ریالی کل', _int(s['total_value'])),
            ('تعداد انواع پالت', _int(s['pallet_type_count'])),
        ],
        'headers': ['ردیف', 'کد پالت', 'نام پالت', 'جنس', 'ابعاد', 'جمع ورود', 'جمع خروج', 'موجودی', 'ارزش ورود (ریال)', 'ارزش خروج (ریال)', 'ارزش فعلی (ریال)', 'آخرین گردش'],
        'rows': rows,
        'widths': [6, 10, 20, 10, 14, 10, 10, 10, 18, 18, 18, 12],
    }])


def export_stock_value_to_pdf(report, file_path) -> str:
    s = report['summary']
    w = _usable_width()
    story = [
        _text_block('گزارش موجودی و ارزش ریالی انبار', 16),
        _text_block(f"انبار:   {report['warehouse']['name']}   |   تاریخ:   {_today_fa()}   |   مبالغ به ریال", 9, '#64748b'),
        Spacer(1, 4),
        _summary_table([
            ('نام انبار', report['warehouse']['name']),
            ('کد انبار', report['warehouse']['code']),
            ('تعداد کل موجودی', _fmt(s['total_qty'])),
            ('ارزش ریالی کل', _money(s['total_value'])),
            ('تعداد انواع پالت', _fmt(s['pallet_type_count'])),
        ], w),
        Spacer(1, 8),
    ]
    headers = ['ردیف', 'کد پالت', 'نام پالت', 'جنس', 'ابعاد', 'ورود', 'خروج', 'موجودی', 'ارزش ورود', 'ارزش خروج', 'ارزش فعلی', 'آخرین گردش']
    rows = []
    for i, it in enumerate(report['items'], 1):
        rows.append([
            i, it['pallet_code'], it['pallet_name'], it['material_type'], it['dimensions'],
            _fmt(it['total_in_qty']), _fmt(it['total_out_qty']), _fmt(it['current_qty']),
            _money(it['total_in_value']), _money(it['total_out_value']), _money(it['current_value']),
            it.get('last_transaction_date') or '-',
        ])
    widths = [1.2, 2.0, 3.0, 1.8, 2.2, 1.9, 1.9, 1.9, 3.0, 3.0, 3.0, 1.9]
    story.append(_data_table(headers, rows, [x * cm for x in widths], font_size=7.5))
    return _build_pdf(file_path, story)


# ================================================================
# 📋 کاردکس
# ================================================================
def export_kardex_to_excel(report, file_path) -> str:
    s = report['summary']
    pallet_name = report['pallet']['name'] if report.get('pallet') else 'همه پالت‌ها'
    rows = []
    for it in report['items']:
        delta = _int(it.get('total_price'))
        if it.get('transaction_type') != 'IN':
            delta = -delta
        rows.append([
            _int(it.get('row_no')), it.get('transaction_date') or '-',
            'ورود' if it.get('transaction_type') == 'IN' else 'خروج',
            it.get('main_reference_no') or '-', it.get('reference_no') or '-',
            it.get('pallet_code') or '-', it.get('pallet_name') or '-',
            _int(it.get('qty_in')), _int(it.get('qty_out')), _int(it.get('running_qty')),
            _int(it.get('unit_price')), delta, _int(it.get('running_value')),
            it.get('description') or '-',
        ])
    return _write_excel(file_path, [{
        'name': 'کاردکس',
        'title': f"کاردکس انبار:   {report['warehouse']['name']}   -   {pallet_name}",
        'subtitle': f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال',
        'summary': [
            ('جمع ورود تعدادی', _int(s['total_in_qty'])),
            ('جمع خروج تعدادی', _int(s['total_out_qty'])),
            ('مانده تعدادی', _int(s['current_qty'])),
            ('مانده ریالی', _int(s['current_value'])),
            ('تعداد تراکنش', _int(s['transaction_count'])),
        ],
        'headers': ['ردیف', 'تاریخ', 'نوع', 'مرجع اصلی', 'سند مرحله', 'کد پالت', 'نام پالت', 'ورود', 'خروج', 'مانده', 'قیمت واحد (ریال)', 'مبلغ گردش (ریال)', 'مانده ریالی (ریال)', 'شرح'],
        'rows': rows,
        'widths': [6, 11, 7, 12, 12, 9, 16, 8, 8, 9, 14, 16, 16, 20],
    }])


def export_kardex_to_pdf(report, file_path) -> str:
    s = report['summary']
    pallet_name = report['pallet']['name'] if report.get('pallet') else 'همه پالت‌ها'
    w = _usable_width()
    story = [
        _text_block('کاردکس پالت و انبار', 16),
        _text_block(f"انبار:   {report['warehouse']['name']}   |   پالت:   {pallet_name}   |   تاریخ:   {_today_fa()}   |   مبالغ به ریال", 9, '#64748b'),
        Spacer(1, 4),
        _summary_table([
            ('جمع ورود تعدادی', _fmt(s['total_in_qty'])),
            ('جمع خروج تعدادی', _fmt(s['total_out_qty'])),
            ('مانده تعدادی', _fmt(s['current_qty'])),
            ('مانده ریالی', _money(s['current_value'])),
            ('تعداد تراکنش', _fmt(s['transaction_count'])),
        ], w),
        Spacer(1, 8),
    ]
    headers = ['ردیف', 'تاریخ', 'نوع', 'مرجع اصلی', 'سند مرحله', 'کد پالت', 'نام پالت', 'ورود', 'خروج', 'مانده', 'قیمت واحد', 'مبلغ گردش', 'مانده ریالی', 'شرح']
    rows = []
    for it in report['items']:
        delta = _int(it.get('total_price'))
        if it.get('transaction_type') != 'IN':
            delta = -delta
        rows.append([
            _int(it.get('row_no')), it.get('transaction_date') or '-',
            'ورود' if it.get('transaction_type') == 'IN' else 'خروج',
            it.get('main_reference_no') or '-', it.get('reference_no') or '-',
            it.get('pallet_code') or '-', it.get('pallet_name') or '-',
            _fmt(it.get('qty_in')), _fmt(it.get('qty_out')), _fmt(it.get('running_qty')),
            _money(it.get('unit_price')), _money(delta), _money(it.get('running_value')),
            it.get('description') or '-',
        ])
    widths = [1.0, 2.2, 1.6, 2.2, 2.0, 1.6, 2.4, 1.6, 1.6, 1.8, 1.8, 3.0, 3.0, 1.2]
    story.append(_data_table(headers, rows, [x * cm for x in widths], font_size=7))
    return _build_pdf(file_path, story)


# ================================================================
# 🗂️ تجمیعی همه انبارها
# ================================================================
def export_aggregate_stock_to_excel(report, file_path) -> str:
    s = report['summary']
    summary = [
        ('تعداد انبارها', _int(s['warehouse_count'])),
        ('تعداد کل موجودی', _int(s['total_qty'])),
        ('ارزش ریالی کل', _int(s['total_value'])),
        ('تعداد انواع پالت', _int(s['pallet_type_count'])),
    ]
    wh_rows = []
    for i, it in enumerate(report['warehouses'], 1):
        wh_rows.append([
            i, it['code'], it['name'], _int(it.get('pallet_type_count')),
            _int(it['total_in_qty']), _int(it['total_out_qty']), _int(it['current_qty']),
            _int(it['total_in_value']), _int(it['total_out_value']), _int(it['current_value']),
            it.get('last_transaction_date') or '-',
        ])
    pl_rows = []
    for i, it in enumerate(report['pallets'], 1):
        pl_rows.append([
            i, it['pallet_code'], it['pallet_name'], it['material_type'], it['dimensions'],
            _int(it['total_in_qty']), _int(it['total_out_qty']), _int(it['current_qty']),
            _int(it['current_value']), it.get('last_transaction_date') or '-',
        ])
    return _write_excel(file_path, [
        {
            'name': 'خلاصه انبارها',
            'title': 'گزارش تجمیعی همه انبارها',
            'subtitle': f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال',
            'summary': summary,
            'headers': ['ردیف', 'کد انبار', 'نام انبار', 'انواع پالت', 'جمع ورود', 'جمع خروج', 'موجودی', 'ارزش ورود (ریال)', 'ارزش خروج (ریال)', 'ارزش فعلی (ریال)', 'آخرین گردش'],
            'rows': wh_rows,
            'widths': [6, 10, 22, 10, 10, 10, 10, 18, 18, 18, 12],
        },
        {
            'name': 'خلاصه پالت‌ها',
            'title': 'خلاصه پالت‌ها در همه انبارها',
            'subtitle': f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال',
            'headers': ['ردیف', 'کد پالت', 'نام پالت', 'جنس', 'ابعاد', 'جمع ورود', 'جمع خروج', 'موجودی', 'ارزش فعلی (ریال)', 'آخرین گردش'],
            'rows': pl_rows,
            'widths': [6, 10, 22, 10, 14, 10, 10, 10, 18, 12],
        },
    ])


def export_aggregate_stock_to_pdf(report, file_path) -> str:
    s = report['summary']
    w = _usable_width()
    story = [
        _text_block('گزارش تجمیعی همه انبارها', 16),
        _text_block(f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال', 9, '#64748b'),
        Spacer(1, 4),
        _summary_table([
            ('تعداد انبارها', _fmt(s['warehouse_count'])),
            ('تعداد کل موجودی', _fmt(s['total_qty'])),
            ('ارزش ریالی کل', _money(s['total_value'])),
            ('تعداد انواع پالت', _fmt(s['pallet_type_count'])),
        ], w),
        Spacer(1, 8),
        _text_block('خلاصه به تفکیک انبار', 12, '#1e40af', 'RIGHT'),
        Spacer(1, 4),
    ]
    wh_headers = ['ردیف', 'کد انبار', 'نام انبار', 'انواع', 'ورود', 'خروج', 'موجودی', 'ارزش ورود', 'ارزش خروج', 'ارزش فعلی', 'آخرین گردش']
    wh_rows = []
    for i, it in enumerate(report['warehouses'], 1):
        wh_rows.append([
            i, it['code'], it['name'], _fmt(it.get('pallet_type_count')),
            _fmt(it['total_in_qty']), _fmt(it['total_out_qty']), _fmt(it['current_qty']),
            _money(it['total_in_value']), _money(it['total_out_value']), _money(it['current_value']),
            it.get('last_transaction_date') or '-',
        ])
    wh_widths = [1.0, 2.0, 3.6, 1.8, 2.2, 2.2, 2.2, 3.2, 3.2, 3.2, 1.9]
    story.append(_data_table(wh_headers, wh_rows, [x * cm for x in wh_widths], font_size=7.5))
    story.append(Spacer(1, 10))
    story.append(_text_block('خلاصه پالت‌ها در همه انبارها', 12, '#1e40af', 'RIGHT'))
    story.append(Spacer(1, 4))

    pl_headers = ['ردیف', 'کد پالت', 'نام پالت', 'جنس', 'ابعاد', 'ورود', 'خروج', 'موجودی', 'ارزش فعلی', 'آخرین گردش']
    pl_rows = []
    for i, it in enumerate(report['pallets'], 1):
        pl_rows.append([
            i, it['pallet_code'], it['pallet_name'], it['material_type'], it['dimensions'],
            _fmt(it['total_in_qty']), _fmt(it['total_out_qty']), _fmt(it['current_qty']),
            _money(it['current_value']), it.get('last_transaction_date') or '-',
        ])
    pl_widths = [1.0, 2.0, 4.2, 2.0, 2.6, 2.2, 2.2, 2.2, 3.6, 2.2]
    story.append(_data_table(pl_headers, pl_rows, [x * cm for x in pl_widths], font_size=7.5))

    return _build_pdf(file_path, story)
# ================================================================
# 🌐 توابع ساخت HTML برای پیش‌نمایش Flask
# ================================================================
import html as _html_lib


def _esc(value: Any) -> str:
    return _html_lib.escape(str(value if value is not None else ''))


def _hnum(value: Any) -> str:
    try:
        return f'{int(value):,}'
    except Exception:
        return str(value) if value is not None else '-'


def _hmoney(value: Any) -> str:
    try:
        return f'{int(value):,} ریال'
    except Exception:
        return '-'


_HTML_STYLE = """
* { box-sizing: border-box; }
body {
    font-family: Tahoma, 'Vazirmatn', 'Segoe UI', sans-serif;
    background: #f1f5f9;
    color: #0f172a;
    margin: 0;
    padding: 24px;
    direction: rtl;
}
.header {
    background: #1e293b;
    color: #fff;
    border-radius: 12px;
    padding: 18px 24px;
    margin-bottom: 18px;
}
.header h1 { margin: 0 0 6px 0; font-size: 20px; }
.header .sub { color: #94a3b8; font-size: 12px; }
.cards { display: flex; flex-wrap: wrap; gap: 12px; margin-bottom: 18px; }
.card {
    background: #fff;
    border: 1px solid #e2e8f0;
    border-radius: 10px;
    padding: 10px 14px;
    min-width: 150px;
    flex: 1;
    box-shadow: 0 1px 3px rgba(0,0,0,0.06);
}
.card .card-title { color: #64748b; font-size: 12px; margin-bottom: 4px; }
.card .card-value { font-size: 20px; font-weight: bold; color: #0f172a; }
.card.green .card-value { color: #059669; }
.card.red .card-value { color: #dc2626; }
.card.blue .card-value { color: #2563eb; }
.card.orange .card-value { color: #ea580c; }
.section-title {
    font-size: 15px; font-weight: bold; color: #1e40af;
    margin: 20px 0 10px 0;
    border-right: 4px solid #1e40af; padding-right: 8px;
}
table {
    width: 100%; border-collapse: collapse; background: #fff;
    border-radius: 10px; overflow: hidden;
    box-shadow: 0 1px 3px rgba(0,0,0,0.08); font-size: 12px;
}
thead th {
    background: #1e40af; color: #fff; padding: 10px 8px;
    text-align: right; font-weight: bold; white-space: nowrap;
}
tbody td { padding: 8px; border-bottom: 1px solid #e2e8f0; text-align: right; }
tbody tr:nth-child(even) { background: #f8fafc; }
tbody tr:hover { background: #eef2ff; }
tr.row-in td { background: #ecfdf5; }
tr.row-in:hover td { background: #d1fae5; }
tr.row-out td { background: #fef2f2; }
tr.row-out:hover td { background: #fee2e2; }
.footer { margin-top: 24px; text-align: center; color: #94a3b8; font-size: 11px; }
"""


def _html_page(title: str, body: str) -> str:
    return (
        '<!DOCTYPE html>\n<html lang="fa" dir="rtl">\n<head>\n<meta charset="utf-8">\n'
        f'<title>{_esc(title)}</title>\n<style>{_HTML_STYLE}</style>\n'
        '</head>\n<body>\n' + body + '\n</body>\n</html>'
    )


def _html_header(title: str, subtitle: str) -> str:
    return (
        '<div class="header">'
        f'<h1>{_esc(title)}</h1>'
        f'<div class="sub">{_esc(subtitle)}</div>'
        '</div>'
    )


def _html_cards(pairs) -> str:
    items = []
    for entry in pairs:
        label, value = entry[0], entry[1]
        kind = entry[2] if len(entry) > 2 else ''
        items.append(
            f'<div class="card {kind}">'
            f'<div class="card-title">{_esc(label)}</div>'
            f'<div class="card-value">{_esc(value)}</div></div>'
        )
    return '<div class="cards">' + ''.join(items) + '</div>'


def _html_table(headers, rows, row_classes=None) -> str:
    ths = ''.join(f'<th>{_esc(h)}</th>' for h in headers)
    body = []
    for i, r in enumerate(rows):
        cls = ''
        if row_classes and i < len(row_classes) and row_classes[i]:
            cls = f' class="{row_classes[i]}"'
        tds = ''.join(f'<td>{_esc(c)}</td>' for c in r)
        body.append(f'<tr{cls}>{tds}</tr>')
    return '<table><thead><tr>' + ths + '</tr></thead><tbody>' + ''.join(body) + '</tbody></table>'


# ----------------------------------------------------------------
# 🏭 ارزش ریالی انبار
# ----------------------------------------------------------------
def build_stock_value_html(report: Dict) -> str:
    s = report['summary']
    wh = report['warehouse']
    parts = [
        _html_header('گزارش موجودی و ارزش ریالی انبار',
                     f"انبار:   {wh['name']}   |   تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال"),
        _html_cards([
            ('نام انبار', wh['name'], 'blue'),
            ('کد انبار', wh.get('code') or '-', ''),
            ('تعداد کل موجودی', f"{_hnum(s['total_qty'])} عدد", 'blue'),
            ('ارزش ریالی کل', _hmoney(s['total_value']), 'green'),
            ('تعداد انواع پالت', _hnum(s['pallet_type_count']), 'orange'),
        ]),
    ]
    headers = ['ردیف', 'کد پالت', 'نام پالت', 'جنس', 'ابعاد', 'جمع ورود', 'جمع خروج',
               'موجودی', 'ارزش ورود', 'ارزش خروج', 'ارزش فعلی', 'آخرین گردش']
    rows = []
    for i, it in enumerate(report['items'], 1):
        rows.append([
            i, it['pallet_code'], it['pallet_name'], it['material_type'], it['dimensions'],
            _hnum(it['total_in_qty']), _hnum(it['total_out_qty']), _hnum(it['current_qty']),
            _hmoney(it['total_in_value']), _hmoney(it['total_out_value']), _hmoney(it['current_value']),
            it.get('last_transaction_date') or '-',
        ])
    parts.append(_html_table(headers, rows))
    parts.append(f'<div class="footer">سیستم مدیریت پالت انبار   |   {_today_fa()}</div>')
    return _html_page('گزارش ارزش ریالی انبار', ''.join(parts))


# ----------------------------------------------------------------
# 📋 کاردکس
# ----------------------------------------------------------------
def build_kardex_html(report: Dict) -> str:
    s = report['summary']
    pallet_name = report['pallet']['name'] if report.get('pallet') else 'همه پالت‌ها'
    parts = [
        _html_header('کاردکس پالت و انبار',
                     f"انبار:   {report['warehouse']['name']}   |   پالت:   {pallet_name}   |   تاریخ:   {_today_fa()}   |   مبالغ به ریال"),
        _html_cards([
            ('جمع ورود تعدادی', _hnum(s['total_in_qty']), 'green'),
            ('جمع خروج تعدادی', _hnum(s['total_out_qty']), 'red'),
            ('مانده تعدادی', _hnum(s['current_qty']), 'blue'),
            ('مانده ریالی', _hmoney(s['current_value']), 'orange'),
            ('تعداد تراکنش', _hnum(s['transaction_count']), ''),
        ]),
    ]
    headers = ['ردیف', 'تاریخ', 'نوع', 'مرجع اصلی', 'سند مرحله', 'کد پالت', 'نام پالت',
               'ورود', 'خروج', 'مانده', 'قیمت واحد', 'مبلغ گردش', 'مانده ریالی', 'شرح']
    rows, classes = [], []
    for it in report['items']:
        delta = _int(it.get('total_price'))
        if it.get('transaction_type') != 'IN':
            delta = -delta
        rows.append([
            _int(it.get('row_no')), it.get('transaction_date') or '-',
            'ورود' if it.get('transaction_type') == 'IN' else 'خروج',
            it.get('main_reference_no') or '-', it.get('reference_no') or '-',
            it.get('pallet_code') or '-', it.get('pallet_name') or '-',
            _hnum(it.get('qty_in')), _hnum(it.get('qty_out')), _hnum(it.get('running_qty')),
            _hmoney(it.get('unit_price')), _hmoney(delta), _hmoney(it.get('running_value')),
            it.get('description') or '-',
        ])
        classes.append('row-in' if it.get('transaction_type') == 'IN' else 'row-out')
    parts.append(_html_table(headers, rows, classes))
    parts.append(f'<div class="footer">سیستم مدیریت پالت انبار   |   {_today_fa()}</div>')
    return _html_page('کاردکس پالت و انبار', ''.join(parts))


# ----------------------------------------------------------------
# 🗂️ تجمیعی همه انبارها
# ----------------------------------------------------------------
def build_aggregate_stock_html(report: Dict) -> str:
    s = report['summary']
    parts = [
        _html_header('گزارش تجمیعی همه انبارها',
                     f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال'),
        _html_cards([
            ('تعداد انبارها', _hnum(s['warehouse_count']), 'blue'),
            ('تعداد کل موجودی', _hnum(s['total_qty']), 'blue'),
            ('ارزش ریالی کل', _hmoney(s['total_value']), 'green'),
            ('تعداد انواع پالت', _hnum(s['pallet_type_count']), 'orange'),
        ]),
        '<div class="section-title">خلاصه به تفکیک انبار</div>',
    ]
    wh_headers = ['ردیف', 'کد انبار', 'نام انبار', 'انواع پالت', 'جمع ورود', 'جمع خروج',
                  'موجودی', 'ارزش ورود', 'ارزش خروج', 'ارزش فعلی', 'آخرین گردش']
    wh_rows = []
    for i, it in enumerate(report['warehouses'], 1):
        wh_rows.append([
            i, it['code'], it['name'], _hnum(it.get('pallet_type_count')),
            _hnum(it['total_in_qty']), _hnum(it['total_out_qty']), _hnum(it['current_qty']),
            _hmoney(it['total_in_value']), _hmoney(it['total_out_value']), _hmoney(it['current_value']),
            it.get('last_transaction_date') or '-',
        ])
    parts.append(_html_table(wh_headers, wh_rows))
    parts.append('<div class="section-title">خلاصه پالت‌ها در همه انبارها</div>')

    pl_headers = ['ردیف', 'کد پالت', 'نام پالت', 'جنس', 'ابعاد', 'جمع ورود', 'جمع خروج',
                  'موجودی', 'ارزش فعلی', 'آخرین گردش']
    pl_rows = []
    for i, it in enumerate(report['pallets'], 1):
        pl_rows.append([
            i, it['pallet_code'], it['pallet_name'], it['material_type'], it['dimensions'],
            _hnum(it['total_in_qty']), _hnum(it['total_out_qty']), _hnum(it['current_qty']),
            _hmoney(it['current_value']), it.get('last_transaction_date') or '-',
        ])
    parts.append(_html_table(pl_headers, pl_rows))
    parts.append(f'<div class="footer">سیستم مدیریت پالت انبار   |   {_today_fa()}</div>')
    return _html_page('گزارش تجمیعی همه انبارها', ''.join(parts))


# ----------------------------------------------------------------
# 📑 اسناد مالی
# ----------------------------------------------------------------
def build_financial_summary_html(report: Dict) -> str:
    s = report['summary']
    parts = [
        _html_header('گزارش اسناد مالی',
                     f'تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال'),
        _html_cards([
            ('تعداد کل اسناد', _hnum(s['total_docs']), 'blue'),
            ('اسناد باز', _hnum(s['open_docs']), 'orange'),
            ('اسناد تسویه', _hnum(s['settled_docs']), 'green'),
            ('مانده دریافتنی', _hmoney(s['receivable_balance']), 'green'),
            ('مانده پرداختنی', _hmoney(s['payable_balance']), 'red'),
        ]),
    ]
    headers = ['شماره سند', 'نوع عملیات', 'جهت', 'طرف حساب', 'مرجع', 'مبلغ کل', 'تسویه شده', 'مانده', 'وضعیت']
    rows = []
    for it in report['items']:
        rows.append([
            it['finance_no'], it['operation_type_label'], it['direction_label'],
            it.get('counterparty_name') or '-', it['reference_label'],
            _hmoney(it.get('total_amount')), _hmoney(it.get('settled_amount')),
            _hmoney(it.get('remaining_amount')), it['status_label'],
        ])
    parts.append(_html_table(headers, rows))
    parts.append(f'<div class="footer">سیستم مدیریت پالت انبار   |   {_today_fa()}</div>')
    return _html_page('گزارش اسناد مالی', ''.join(parts))


# ----------------------------------------------------------------
# 👥 گردش حساب شخص
# ----------------------------------------------------------------
def build_person_statement_html(statement: Dict) -> str:
    person = statement['person']
    summary = statement['summary']
    full_name = f"{person['first_name']} {person['last_name']}".strip()
    parts = [
        _html_header(f'گردش حساب:   {full_name}',
                     f"موبایل:   {person.get('mobile') or '-'}   |   تاریخ گزارش:   {_today_fa()}   |   مبالغ به ریال"),
        _html_cards([
            ('کل دریافتنی', _hmoney(summary['receivable_total']), 'green'),
            ('کل پرداختنی', _hmoney(summary['payable_total']), 'red'),
            ('خالص مانده', _hmoney(summary['net_balance']), 'blue'),
        ]),
    ]
    headers = ['تاریخ', 'نوع ردیف', 'شماره سند', 'نوع عملیات', 'شرح', 'بدهکار', 'بستانکار', 'مانده']
    rows = []
    for d in statement['rows']:
        running = _int(d['running_balance'])
        running_text = f'{_hmoney(abs(running))}  {"بدهکار" if running >= 0 else "بستانکار"}'
        rows.append([
            _fa_date(d['event_date']), ROW_TYPE_LABELS.get(d['row_type'], d['row_type']),
            d['reference_no'], OPERATION_LABELS.get(d['operation_type'], d['operation_type']),
            d.get('description') or '-',
            _hmoney(d['debit_amount']) if _int(d['debit_amount']) else '-',
            _hmoney(d['credit_amount']) if _int(d['credit_amount']) else '-',
            running_text,
        ])
    parts.append(_html_table(headers, rows))
    parts.append(f'<div class="footer">سیستم مدیریت پالت انبار   |   {_today_fa()}</div>')
    return _html_page(f'گردش حساب {full_name}', ''.join(parts))

# ================================================================
# 🏦 گزارش PDF حساب‌های بانکی شخص
# ================================================================
def export_person_bank_accounts_to_pdf(data: Dict, file_path) -> str:
    accounts = data.get('accounts') or []
    w = _usable_width()
    default_bank = next((a.get('bank_name') or '-' for a in accounts if a.get('is_default')), '-')
    story = [
        _text_block('گزارش حساب‌های بانکی شخص', 16),
        _text_block(
            f"شخص:   {data.get('person_name') or '-'}   |   "
            f"موبایل:   {data.get('person_mobile') or '-'}   |   تاریخ:   {_today_fa()}",
            9, '#64748b',
        ),
        Spacer(1, 4),
        _summary_table([
            ('تعداد حساب‌ها', _fmt(len(accounts))),
            ('حساب پیش‌فرض', default_bank),
        ], w),
        Spacer(1, 8),
    ]
    headers = ['ردیف', 'نام بانک', 'شماره حساب', 'شماره شبا', 'شماره کارت', 'نام شعبه', 'کد شعبه', 'پیش‌فرض']
    rows = []
    for i, a in enumerate(accounts, 1):
        rows.append([
            i, a.get('bank_name') or '-', a.get('account_number') or '-',
            a.get('iban') or '-', a.get('card_number') or '-',
            a.get('branch_name') or '-', a.get('branch_code') or '-',
            'بله' if a.get('is_default') else 'خیر',
        ])
    widths = [1.2, 4.0, 3.4, 5.4, 4.2, 3.2, 2.4, 1.8]
    story.append(_data_table(headers, rows, [x * cm for x in widths], font_size=8))
    return _build_pdf(file_path, story)